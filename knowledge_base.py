from __future__ import annotations

from functools import lru_cache
from pathlib import Path
import re

from docx import Document
from langchain_core.tools import tool
from pypdf import PdfReader

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DOCX_PATH = DATA_DIR / "Kurikulum_Profil_Prodi_TI_UIN_Palangka_Raya.docx"
PDF_PATH = DATA_DIR / "Kepmen_553_B_O_2026_Izin_Prodi_TI.pdf"

SOURCE_DOCX = "Kurikulum dan Profil Prodi Teknologi Informasi UIN Palangka Raya"
SOURCE_PDF = "Keputusan Menteri Pendidikan Tinggi, Sains, dan Teknologi Nomor 553/B/O/2026"

# Ringkasan terstruktur untuk tool yang paling sering dipakai.
GENERAL_INFO = """
Program Studi Teknologi Informasi UIN Palangka Raya merupakan Program Sarjana yang dirancang
untuk membekali mahasiswa agar mampu memahami, merancang, mengimplementasikan, dan mengelola
teknologi informasi untuk mendukung kebutuhan organisasi dan masyarakat.

Visi keilmuan:
“Menjadi Program Studi Teknologi Informasi yang unggul di Kalimantan pada tahun 2030 dengan
integrasi nilai-nilai Islam serta berkontribusi pada transformasi digital dan pengelolaan
lingkungan melalui pengembangan infrastruktur teknologi informasi, sistem terintegrasi,
dan teknologi cerdas berbasis jaringan dan komputasi awan.”

Distingsi utama:
1. Fokus pada pengelolaan infrastruktur teknologi informasi dan integrasi sistem.
2. Penerapan cloud computing, jaringan komputer, dan keamanan sistem.
3. Pengembangan Internet of Things (IoT) untuk smart environment dan smart agriculture.
4. Integrasi nilai-nilai keislaman dalam pengembangan dan pemanfaatan teknologi.
5. Orientasi pada permasalahan lokal Kalimantan, seperti lingkungan, lahan gambut, dan pertanian.
"""

PROFILES = """
Profil lulusan Prodi Teknologi Informasi:
PL01 — Environmental IT Analyst dan Smart Agriculture System Analyst: menganalisis permasalahan
komputasi dan kebutuhan sistem TI pada bidang lingkungan, pertanian digital, dan pengelolaan SDA.
PL02 — Cloud & Infrastructure Engineer berbasis lingkungan dan IoT Engineer untuk monitoring
lingkungan: merancang, mengimplementasikan, mengintegrasikan, dan mengevaluasi sistem berbasis
jaringan, cloud computing, dan IoT untuk monitoring serta pengelolaan lingkungan.
PL03 — IT Professional berbasis nilai keislaman: profesional, beretika, mengintegrasikan nilai
keislaman dalam pengembangan/pemanfaatan TI, serta mampu bekerja dalam tim multidisiplin.
PL04 — Technopreneur berbasis teknologi lingkungan: berpikir logis, kritis, sistematis, dan
mengembangkan inovasi serta kewirausahaan digital berbasis teknologi untuk masyarakat dan lingkungan.

Bidang pekerjaan yang juga disebut dalam dokumen antara lain System Analyst, Network Administrator,
Cloud Engineer, Cyber Security Analyst, Database Administrator, IT Infrastructure Engineer, IoT Engineer,
UI/UX Designer, dan Technopreneur.
"""

CONCENTRATIONS = """
Konsentrasi keilmuan Prodi Teknologi Informasi:
1. Infrastruktur TI & Cloud Computing
   - Virtualisasi Server
   - Administrasi Sistem Linux
   - Cloud Infrastructure
   - DevOps dan Container
   - Cloud Security
   - Proyek Infrastruktur Cloud
2. Sistem Terintegrasi & Transformasi Digital
   - Enterprise System Architecture
   - Business Process Management
   - Integrasi Sistem Informasi
   - Digital Platform Development
   - Smart Government System
   - Proyek Transformasi Digital
3. Smart Environment & IoT
   - Sensor dan Embedded System
   - IoT Monitoring Lingkungan
   - Smart Agriculture Technology
   - Sistem Monitoring Kebakaran Hutan
   - Data Lingkungan Berbasis IoT
   - Proyek Smart Environment
"""

CURRICULUM = """
Dokumen memuat struktur kurikulum utama dan bagian struktur alternatif/MBKM. Pada struktur utama,
total kurikulum dinyatakan 144 SKS.

Ringkasan struktur utama:
Semester 1 — 19 SKS: Studi Agama Kontemporer, Pancasila, Bahasa Indonesia, Kalkulus,
Matematika Diskrit, Algoritma dan Pemrograman, Pengantar Teknologi Informasi, Sistem Digital.
Semester 2 — 22 SKS: Bahasa Inggris, Kewarganegaraan, Struktur Data, Organisasi dan Arsitektur
Komputer, Sistem Operasi, Komunikasi Data, Manajemen Basis Data, Statistika Komputasi.
Semester 3 — 21 SKS: Jaringan Komputer, Pemrograman Web, Pemrograman Mobile, Analisis dan Desain
Sistem, Rekayasa Perangkat Lunak, Interaksi Manusia dan Komputer, Metode Numerik, Bimbingan Membaca Al-Qur’an.
Semester 4 — 21 SKS: Infrastruktur Teknologi Informasi, Keamanan Sistem Informasi, Internet of Things,
Cloud Computing, Manajemen Proyek TI, Tata Kelola TI, Kecerdasan Buatan, Pengembangan Sistem Terintegrasi.
Semester 5 — 21 SKS: Big Data dan Analitik Data, Etika Profesi TI, Metodologi Penelitian, Proyek TI,
Mata Kuliah Konsentrasi 1–2, Praktik Pengamalan Ibadah, Integrasi Islam dan Teknologi.
Semester 6 — 20 SKS pada tabel struktur utama: Magang Industri, Administrasi Sistem dan Server,
Mata Kuliah Konsentrasi 3–5, Publikasi Ilmiah, Infrastruktur Jaringan Smart Agriculture.
Semester 7 — 20 SKS: Mata Kuliah Konsentrasi 6, KKN, Proyek Sistem Informasi Agroekologi,
Proyek IoT Monitoring Lingkungan, Cloud Computing untuk Sistem Lingkungan, Skripsi.

Dokumen juga memuat rincian mata kuliah kekhasan Prodi dan struktur alternatif pada bagian MBKM.
Jika terdapat perbedaan angka SKS antarbagian, chatbot harus menyebutkan perbedaan tersebut dan
mengutamakan konteks pertanyaan pengguna, bukan mengarang angka baru.
"""

CPL_INFO = """
Capaian Pembelajaran Lulusan (CPL) Prodi Teknologi Informasi terdiri dari 9 CPL utama:
CPL01 — mempunyai pengetahuan matematika, computing, dan ilmu lain yang relevan untuk menyelesaikan
permasalahan computing yang kompleks dengan pendekatan teknologi informasi.
CPL02 — mampu menganalisis, mengidentifikasi, dan mendefinisikan permasalahan computing yang kompleks
serta memberikan pendekatan teknologi informasi sebagai solusi.
CPL03 — mampu merancang, mengimplementasikan, dan mengevaluasi solusi berbasis computing sesuai
kebutuhan organisasi dan masyarakat dengan memanfaatkan teknologi informasi modern.
CPL04 — mampu memilih, mengintegrasikan, dan mengadministrasikan infrastruktur TI sebagai solusi.
CPL05 — mampu mengimplementasikan, mengelola, dan mengamankan sistem dan informasi terdistribusi.
CPL06 — mampu berkomunikasi efektif secara lisan maupun tulisan dalam lingkungan profesional TI.
CPL07 — mampu bekerja efektif sebagai individu maupun tim multidisiplin dalam pengembangan dan
pengelolaan sistem TI.
CPL08 — memiliki sikap profesional, etika, tanggung jawab sosial serta memahami dampak TI terhadap
masyarakat dan lingkungan.
CPL09 — memiliki kemampuan belajar sepanjang hayat serta mengikuti perkembangan TI yang cepat.
"""

MBKM_INFO = """
Dokumen merancang MBKM untuk memberikan fleksibilitas pembelajaran di luar program studi/perguruan
tinggi dan memperkuat kompetensi teknis, profesional, etika, serta integritas.

Bentuk kegiatan yang disebut antara lain Magang/Praktik Industri, Studi Independen, Pertukaran Pelajar,
Penelitian/Riset Terapan, KKN Tematik/Smart Village, Wirausaha Digital, dan Proyek Independen.
Mekanisme rekognisi/konversi meliputi pemetaan CPL/CPMK, penilaian kesesuaian aktivitas, penetapan
mata kuliah konversi, skema konversi SKS, serta validasi dan penetapan nilai.
"""

IZIN_PRODI = """
Keputusan Menteri Pendidikan Tinggi, Sains, dan Teknologi Republik Indonesia Nomor 553/B/O/2026
Tentang Izin Pembukaan Program Studi Teknologi Informasi Program Sarjana pada Universitas Islam
Negeri Palangka Raya di Kota Palangka Raya yang diselenggarakan oleh Kementerian Agama.

Ditetapkan di Jakarta pada 18 Mei 2026.
Diktum KESATU: memberikan izin pembukaan Program Studi Teknologi Informasi Program Sarjana pada
Universitas Islam Negeri Palangka Raya di Kota Palangka Raya yang diselenggarakan oleh Kementerian Agama.
Diktum KEDUA: Program Studi sebagaimana dimaksud dalam Diktum KESATU dinyatakan memenuhi persyaratan
minimum akreditasi.
Diktum KETIGA: UIN Palangka Raya wajib memenuhi standar nasional pendidikan tinggi dan melaporkan
hasil penyelenggaraan program studi paling lambat 1 bulan setelah akhir setiap semester kepada Menteri.
"""


def _clean(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_docx() -> str:
    if not DOCX_PATH.exists():
        return ""
    doc = Document(DOCX_PATH)
    parts: list[str] = []
    # Paragraf
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())
    # Tabel: seluruh isi sel ikut dimasukkan agar informasi mata kuliah/CPL tidak hilang.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " | ") for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return _clean("\n".join(parts))


def _extract_pdf() -> str:
    if not PDF_PATH.exists():
        return ""
    reader = PdfReader(str(PDF_PATH))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        txt = page.extract_text() or ""
        if txt.strip():
            pages.append(f"[Halaman {i}]\n{txt.strip()}")
    return _clean("\n\n".join(pages))


@lru_cache(maxsize=1)
def load_source_documents() -> dict[str, str]:
    return {SOURCE_DOCX: _extract_docx(), SOURCE_PDF: _extract_pdf()}


def _chunks(text: str, size: int = 1100, overlap: int = 150):
    text = _clean(text)
    if not text:
        return []
    out = []
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        if end < len(text):
            cut = max(text.rfind("\n", start, end), text.rfind(". ", start, end))
            if cut > start + size // 2:
                end = cut + (1 if text[cut] == "\n" else 2)
        out.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return out


@lru_cache(maxsize=1)
def build_search_index():
    index = []
    for source, text in load_source_documents().items():
        for chunk in _chunks(text):
            index.append((source, chunk))
    return index


def _search(query: str, limit: int = 5) -> str:
    q = query.lower().strip()
    terms = [t for t in re.findall(r"[\w&-]+", q) if len(t) >= 3]
    if not terms:
        return "Masukkan kata kunci pencarian."

    scored = []
    for source, chunk in build_search_index():
        low = chunk.lower()
        score = sum(low.count(term) for term in terms)
        # Bonus untuk frasa utuh.
        if q in low:
            score += 8
        if score:
            scored.append((score, source, chunk))
    scored.sort(key=lambda x: x[0], reverse=True)

    if not scored:
        return "Tidak ditemukan informasi yang relevan dalam dua dokumen sumber."

    results = []
    for i, (_, source, chunk) in enumerate(scored[:limit], start=1):
        results.append(f"[{i}] Sumber: {source}\n{chunk}")
    return "\n\n".join(results)


@tool
def get_general_info() -> str:
    """Mengambil informasi umum, visi, keunggulan, dan distingsi Program Studi Teknologi Informasi UIN Palangka Raya."""
    return GENERAL_INFO


@tool
def get_profil_lulusan() -> str:
    """Mengambil profil lulusan dan bidang pekerjaan Program Studi Teknologi Informasi UIN Palangka Raya."""
    return PROFILES


@tool
def get_kurikulum() -> str:
    """Mengambil struktur kurikulum dan mata kuliah Program Studi Teknologi Informasi."""
    return CURRICULUM


@tool
def get_konsentrasi() -> str:
    """Mengambil tiga konsentrasi keilmuan dan mata kuliah konsentrasi Prodi Teknologi Informasi."""
    return CONCENTRATIONS


@tool
def get_cpl() -> str:
    """Mengambil sembilan Capaian Pembelajaran Lulusan utama Prodi Teknologi Informasi."""
    return CPL_INFO


@tool
def get_mbkm() -> str:
    """Mengambil informasi kegiatan dan mekanisme MBKM Prodi Teknologi Informasi."""
    return MBKM_INFO


@tool
def get_izin_prodi() -> str:
    """Mengambil status izin pembukaan Program Studi Teknologi Informasi berdasarkan Keputusan Menteri Nomor 553/B/O/2026."""
    return IZIN_PRODI


@tool
def search_information(query: str) -> str:
    """Mencari informasi faktual dari dua dokumen sumber Prodi TI berdasarkan kata kunci pengguna."""
    return _search(query)
